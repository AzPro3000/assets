#!/usr/bin/env python3
"""doc-informe · Maqueta un informe de Azolla Projects con la identidad V3 (IBM Plex Sans, clean, A4).

Uso:
    python3 maquetar_informe.py contenido.txt salida.docx [--pdf salida.pdf] [--sin-render]

contenido.txt usa el formato descrito en references/formato-contenido.md.
El script:
  1. construye el .docx (estilos editables, tablas reales, cabecera, pie y helecho);
  2. si hay LibreOffice, lo renderiza, detecta los capítulos que caen en cabeza de página
     y les quita la línea separadora (evita la doble línea con la cabecera); repite hasta estabilizar;
  3. comprueba palabra a palabra que el texto del .docx es el del contenido.txt.
"""
import re, sys, os, json, shutil, subprocess, tempfile, argparse
from PIL import Image
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = "https://raw.githubusercontent.com/AzPro3000/assets/main"
FONTS_URL = "https://raw.githubusercontent.com/IBM/plex/master/packages/plex-sans/fonts/complete/ttf/"
ASSETS = os.environ.get("DOC_INFORME_ASSETS") or (
    os.path.join(HERE, "..", "assets") if os.path.isdir(os.path.join(HERE, "..", "assets"))
    else os.path.join(tempfile.gettempdir(), "doc-informe-assets"))

def ensure_assets():
    """Recursos de marca. Nunca se redibujan: se descargan de los originales.
    - Fuentes IBM Plex Sans (Regular, Medium, SemiBold) del repositorio oficial de IBM.
    - Logo monocromo: silueta (canal alfa) del logo oficial coloreada en negro cálido.
    - Helecho: recurso original aprobado (doc-informe/helecho.png en el repositorio de Azolla Projects)."""
    import urllib.request
    os.makedirs(os.path.join(ASSETS, "fonts"), exist_ok=True)
    for w in ("Regular", "Medium", "SemiBold"):
        p = os.path.join(ASSETS, "fonts", "IBMPlexSans-%s.ttf" % w)
        if not os.path.exists(p): urllib.request.urlretrieve(FONTS_URL + "IBMPlexSans-%s.ttf" % w, p)
    logo = os.path.join(ASSETS, "logo_mono.png")
    if not os.path.exists(logo):
        src = os.path.join(ASSETS, "logo_dark_text.png")
        if not os.path.exists(src): urllib.request.urlretrieve(REPO + "/logo_dark_text.png", src)
        im = Image.open(src).convert("RGBA"); im = im.crop(im.getchannel("A").getbbox())
        out = Image.new("RGBA", im.size, (0x22, 0x20, 0x1C, 0)); out.putalpha(im.getchannel("A"))
        out.resize((700, round(700 * im.height / im.width)), Image.LANCZOS).save(logo, optimize=True)
    fern = os.path.join(ASSETS, "helecho.png")
    if not os.path.exists(fern):
        try: urllib.request.urlretrieve(REPO + "/doc-informe/helecho.png", fern)
        except Exception as e:
            sys.exit("Falta el helecho original (%s). Pide al usuario 01-informe-aprobado.html y extráelo "
                     "con: python3 maquetar_informe.py --extraer-helecho 01-informe-aprobado.html" % e)

# ---------------- identidad V3
OLIVE, VIVID, PALE, INK, BODY, GREY, SEP, SEP2, ALT = (
    "4B7831", "69AF40", "EEF4E8", "22201C", "4A4741", "77746D", "B9C3B2", "DEDED9", "F7F9F5")
F_REG, F_MED, F_SEMI = "IBM Plex Sans", "IBM Plex Sans Medium", "IBM Plex Sans SemiBold"
PAGE_W, PAGE_H = 210.0, 297.0
ML = MR = 15.2
TEXT_W = PAGE_W - ML - MR            # 179,6 mm de ancho útil
COL = 18.0                            # columna de numeración de capítulos
PROSE_R = 15.0                        # la prosa no llega al margen derecho (línea más legible)
LINE_Y_MM = 280.5                     # línea del pie (medida en render con estos márgenes)
KNOT = (0.500, 0.594)                 # nudo de las tres ramas del helecho, en fracción del recurso
LINE = 1.15                           # interlineado múltiplo (≈1,5 em en Plex). Nunca exacto: Google Docs lo rompe
TW = lambda mm: int(round(mm * 56.6929))
CHAR_MM = 1.75                        # ancho medio de carácter a 9,5 pt (para medir columnas)
CELL_PAD = 5.0                        # relleno horizontal total de una celda

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e

def set_font(r, font=F_REG, size=None, color=None):
    r.font.name = font
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = el("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)

def pfmt(p, before=0, after=0, line=None, left=0, right=0, hanging=0, keep_next=False, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line: pf.line_spacing = line           # float -> múltiplo (auto)
    pf.left_indent = Mm(left); pf.right_indent = Mm(right)
    if hanging: pf.first_line_indent = -Mm(hanging)
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if align is not None: pf.alignment = align

def border(p, side, color, sz=6, space=0):
    ppr = p._p.get_or_add_pPr()
    b = ppr.find(qn("w:pBdr"))
    if b is None:
        b = el("w:pBdr"); anchor = None
        for t in ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:widowControl", "w:numPr"):
            x = ppr.find(qn(t))
            if x is not None: anchor = x
        if anchor is not None: anchor.addnext(b)
        else: ppr.insert(0, b)
    b.append(el("w:" + side, **{"w:val": "single", "w:sz": sz, "w:space": space, "w:color": color}))

def add_runs(p, text, size, color=BODY, bold_color=INK):
    """**negrita** -> IBM Plex Sans Medium en negro cálido. '\\n' literal -> salto de línea."""
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part: continue
        bold = part.startswith("**") and part.endswith("**")
        txt = part[2:-2] if bold else part
        pieces = txt.split("\\n")
        for j, piece in enumerate(pieces):
            r = p.add_run(piece)
            set_font(r, F_MED if bold else F_REG, size, bold_color if bold else color)
            if j < len(pieces) - 1: r._element.append(el("w:br"))

def field(p, code, size, color):
    def mk(inner):
        r = p.add_run(); set_font(r, F_REG, size, color); r._element.append(inner)
    mk(el("w:fldChar", **{"w:fldCharType": "begin"}))
    it = el("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " %s " % code; mk(it)
    mk(el("w:fldChar", **{"w:fldCharType": "separate"}))
    r = p.add_run("1"); set_font(r, F_REG, size, color)
    mk(el("w:fldChar", **{"w:fldCharType": "end"}))

def split_cells(line):
    return [c.replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]

# ---------------- lectura del contenido
def parse(path):
    P = {"titulo": [], "para": "", "fecha": "", "cabecera": "", "pie": "", "salto_primer_capitulo": "auto"}
    meta, blocks = [], []
    lines = [l.rstrip("\n") for l in open(path, encoding="utf-8")]
    i = 0
    while i < len(lines):
        l = lines[i]
        if not l.strip(): i += 1; continue
        if l.startswith("@"):
            k, *v = split_cells(l[1:])
            P[k] = v if k == "titulo" else (v[0] if v else "")
        elif l.startswith("META|"):
            _, k, v = (split_cells(l) + ["", ""])[:3]; meta.append((k, v))
        elif l.startswith("# "): blocks.append(("h1", l[2:]))
        elif l.startswith("## "): blocks.append(("h2", l[3:]))
        elif l.startswith("- "): blocks.append(("li", l[2:]))
        elif l.startswith("TABLE"):
            opts = l.split()[1:]; rows = []; widths = None; i += 1
            while lines[i] != "ENDTABLE":
                if lines[i].startswith("WIDTHS|"): widths = [float(x) for x in lines[i].split("|")[1:]]
                elif lines[i].strip(): rows.append(split_cells(lines[i]))
                i += 1
            blocks.append(("table", {"rows": rows, "header": "nohead" not in opts, "widths": widths}))
        else: blocks.append(("p", l))
        i += 1
    if not P["titulo"]:
        sys.exit("Falta @titulo en el contenido.")
    return P, meta, blocks

def plain_words(path):
    """Texto del contenido sin marcas, para la verificación palabra a palabra."""
    out = []
    for l in open(path, encoding="utf-8"):
        l = l.rstrip("\n")
        if not l.strip() or l.startswith("@") or l.startswith("WIDTHS|") or l.startswith("TABLE") or l == "ENDTABLE":
            continue
        l = re.sub(r"^META\|", "", l); l = re.sub(r"^(#+ |- )", "", l)
        l = l.replace("**", "").replace("\\n", " ")
        out += [c.replace("\\|", "|") for c in re.split(r"(?<!\\)\|", l)]
    return " ".join(out).split()

# ---------------- anchos de columna según el contenido
_FONTS = {}
def _font(bold, size):
    from PIL import ImageFont
    key = (bold, size)
    if key not in _FONTS:
        f = os.path.join(ASSETS, "fonts", "IBMPlexSans-Medium.ttf" if bold else "IBMPlexSans-Regular.ttf")
        _FONTS[key] = ImageFont.truetype(f, size * 10)        # medido a 10x para precisión
    return _FONTS[key]

from functools import lru_cache
@lru_cache(maxsize=None)
def text_mm(t, bold=False, size=9.5):
    return _font(bold, size).getlength(t) / 10 * 25.4 / 72

@lru_cache(maxsize=None)
def cell_lines(text, width_mm, bold=False, size=9.5):
    """Líneas que ocupa un texto en una celda de ese ancho (ajuste greedy por palabras)."""
    avail = width_mm - CELL_PAD - 2.0          # margen: Word, LibreOffice y Google Docs no componen igual
    total = 0
    for para in text.replace("**", "").split("\\n"):
        words = para.split() or [""]
        n, cur = 1, 0.0
        sp = text_mm(" ", bold, size)
        for w in words:
            wl = text_mm(w, bold, size)
            if cur == 0: cur = wl
            elif cur + sp + wl <= avail: cur += sp + wl
            else: n += 1; cur = wl
            if wl > avail: n += int(wl // max(avail, 1))     # palabra más ancha que la celda: se partiría
        total += n
    return total

def auto_widths(rows, total, header=True):
    """Reparte el ancho para que la tabla ocupe la menor altura posible.
    Parte del mínimo de cada columna (su palabra más larga) y va dando ancho a la columna
    que más líneas ahorra por milímetro. Así los códigos, fechas y cifras quedan compactos
    (en cuanto caben en una línea dejan de pedir ancho) y las descripciones y acciones,
    que son las que más líneas ahorran, se quedan con el espacio."""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    hdr = [header and i == 0 for i in range(len(rows))]
    size = lambda i: 9.0 if hdr[i] else 9.5
    minw = []
    for c in range(ncol):
        lw = max((text_mm(w, hdr[i], size(i)) for i, r in enumerate(rows) for w in r[c].replace("**", "").replace("\\n", " ").split()), default=5)
        minw.append(max(lw + CELL_PAD + 1.0, 9.0))
    widths = minw[:]
    if sum(widths) > total:      # ni siquiera caben las palabras más largas: se reparte en proporción
        return [w * total / sum(widths) for w in widths]
    def height(ws):
        # altura de la tabla (filas) y, como desempate, el total de líneas de todas las celdas:
        # así también progresa cuando dos celdas de una misma fila tienen que estrecharse a la vez
        lines = [[cell_lines(r[c], ws[c], hdr[i], size(i)) for c in range(ncol)] for i, r in enumerate(rows)]
        return sum(max(l) for l in lines) * 1000 + sum(map(sum, lines))
    widths = [round(w * 2) / 2 for w in widths]
    remaining = total - sum(widths); step = 0.5
    h = height(widths)
    while remaining >= step:
        best, best_gain = None, 0
        for c in range(ncol):
            for k in range(1, int(min(remaining, 60) / step) + 1):
                ws = widths[:]; ws[c] += k * step
                nh = height(ws)
                if nh < h:
                    gain = (h - nh) / (k * step)
                    if gain > best_gain: best, best_gain = (c, k * step, nh), gain
                    break
        if not best: break
        c, add, nh = best
        widths[c] += add; remaining -= add; h = nh
    if remaining > 0:            # ya no se ahorran líneas: el sobrante va a las columnas de más texto
        txt = [sum(text_mm(r[c].replace("**", "")) for i, r in enumerate(rows) if not hdr[i]) + 1 for c in range(ncol)]
        textual = [c for c in range(ncol) if not numeric_col(rows, c, header)] or list(range(ncol))
        long_cols = [c for c in textual if cell_lines(max((r[c] for r in rows), key=len), widths[c]) > 1] or textual
        tot = sum(txt[c] for c in long_cols)
        for c in long_cols: widths[c] += remaining * txt[c] / tot
    return widths

def numeric_col(rows, c, header=True):
    """Columna de cifras comparables (importes, superficies, porcentajes): se alinea a la derecha.
    Índices cortos (1, 2, 3…), códigos y fechas no cuentan: van a la izquierda."""
    body = [r[c].replace("**", "").strip() for r in (rows[1:] if header else rows) if c < len(r) and r[c].strip()]
    if not body: return False
    pat = r"[+−\-–]?[\d.,\s]*\d(?:[.,]\d+)?\s*(%|€|ha|t|t C/ha|kg|m|mm|cm)?(\s*±\s*[\d.,]+)?"
    return all(re.fullmatch(pat, x) for x in body) and any(re.search(r"\d[.,]\d", x) or len(re.sub(r"\D", "", x)) >= 4 for x in body)

# ---------------- construcción
def build(src, out, top_h1=frozenset()):
    P, meta, blocks = parse(src)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(PAGE_W), Mm(PAGE_H)
    sec.left_margin, sec.right_margin = Mm(ML), Mm(MR)
    sec.top_margin, sec.bottom_margin = Mm(38), Mm(31)
    sec.header_distance, sec.footer_distance = Mm(12), Mm(8.5)
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = F_REG; normal.font.size = Pt(10); normal.font.color.rgb = RGBColor.from_string(BODY)
    rpr = normal.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    for a in list(rf.attrib): del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), F_REG)
    rpr.append(el("w:lang", **{"w:val": "es-ES", "w:eastAsia": "es-ES", "w:bidi": "ar-SA"}))
    normal.paragraph_format.space_after = Pt(0); normal.paragraph_format.line_spacing = 1.0
    drf = st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:rFonts"))
    if drf is not None:
        for a in list(drf.attrib): del drf.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): drf.set(qn(a), F_REG)

    def style_heading(name, size, color, font=F_MED):
        s = st[name]
        s.font.name = font; s.font.size = Pt(size); s.font.bold = False; s.font.italic = False
        s.font.color.rgb = RGBColor.from_string(color)
        r = s.element.get_or_add_rPr(); f = r.find(qn("w:rFonts"))
        if f is None: f = el("w:rFonts"); r.insert(0, f)
        for a in list(f.attrib): del f.attrib[a]
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): f.set(qn(a), font)
        for sp in r.findall(qn("w:spacing")): r.remove(sp)       # sin tracking heredado
        pPr = s.element.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:pBdr")) is not None: pPr.remove(pPr.find(qn("w:pBdr")))
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(0); s.paragraph_format.space_after = Pt(0)
    style_heading("Title", 30, INK, F_SEMI)
    style_heading("Heading 1", 16, OLIVE)
    style_heading("Heading 2", 12, INK)

    # cabecera: logo monocromo + rótulo corto, línea fina
    hp = sec.header.paragraphs[0]; hp.style = st["Normal"]; pfmt(hp)
    hp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    hp.add_run().add_picture(os.path.join(ASSETS, "logo_mono.png"), width=Mm(41.0))
    r = hp.add_run("\t" + P["cabecera"]); set_font(r, F_REG, 8.5, GREY)
    r._element.get_or_add_rPr().append(el("w:position", **{"w:val": 7}))
    border(hp, "bottom", SEP, sz=6, space=16)

    # helecho: recurso original recortado a su parte visible, 14 % de opacidad, detrás del texto,
    # con el nudo de las tres ramas sobre la línea del pie (9 mm del borde derecho)
    fern = Image.open(os.path.join(ASSETS, "helecho.png")).convert("RGBA")
    FW = 73.5; FH = FW * fern.height / fern.width
    left = PAGE_W - 9.0 - KNOT[0] * FW
    top = LINE_Y_MM - KNOT[1] * FH
    vis_w, vis_h = PAGE_W - left, min(PAGE_H - top, FH)
    crop = fern.crop((0, 0, round(fern.width * vis_w / FW), round(fern.height * vis_h / FH)))
    crop.putalpha(crop.getchannel("A").point(lambda v: int(v * 0.14)))
    crop = crop.resize((620, round(crop.height * 620 / crop.width)), Image.LANCZOS)
    tmp_fern = os.path.join(tempfile.gettempdir(), "doc_informe_helecho.png"); crop.save(tmp_fern, optimize=True)
    fr = hp.add_run(); fr.add_picture(tmp_fern, width=Mm(vis_w), height=Mm(vis_h))
    inline = fr._element.find(".//" + qn("wp:inline")); ext = inline.find(qn("wp:extent"))
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="1" locked="1" '
        'layoutInCell="0" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionV>'
        '<wp:extent cx="%s" cy="%s"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '<wp:docPr id="900" name="Helecho"/><wp:cNvGraphicFramePr/></wp:anchor>'
        % (Mm(PAGE_W - vis_w), Mm(PAGE_H - vis_h), ext.get("cx"), ext.get("cy")))
    anchor.append(inline.find(qn("a:graphic")))
    inline.getparent().replace(inline, anchor)

    # pie: línea + "Azolla Projects" + rótulo / página
    fp = sec.footer.paragraphs[0]; fp.style = st["Normal"]; pfmt(fp)
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    border(fp, "top", SEP, sz=6, space=10)
    r = fp.add_run("Azolla Projects"); set_font(r, F_REG, 8, GREY)
    r = fp.add_run("\t" + (P["pie"] + "   /   " if P["pie"] else "")); set_font(r, F_REG, 8, GREY)
    field(fp, "PAGE", 8, GREY)

    # apertura: título SemiBold 30 pt; con dos tramos, el primero en negro cálido y el segundo en oliva
    p = doc.add_paragraph(style="Title"); pfmt(p, keep_next=True, line=1.0)
    segs = P["titulo"]
    for j, s in enumerate(segs):
        last = j == len(segs) - 1
        color = OLIVE if (len(segs) > 1 and last) else INK
        r = p.add_run(s + ("" if last else " ")); set_font(r, F_SEMI, 30, color)
        r._element.get_or_add_rPr().append(el("w:spacing", **{"w:val": -16}))
        if not last: r._element.append(el("w:br"))
    if P["para"] or P["fecha"]:
        p = doc.add_paragraph(); pfmt(p, before=14, after=24, keep_next=True)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(P["para"] + "\t" + P["fecha"]); set_font(r, F_REG, 10, GREY)
    else:
        p.paragraph_format.space_after = Pt(24)

    def cell_set(c, width_mm, fill=None, vpad=3.6, lr=2.5, borders=None):
        tcPr = c._tc.get_or_add_tcPr()
        for t in ("w:tcW", "w:shd", "w:tcMar", "w:tcBorders", "w:vAlign"):
            x = tcPr.find(qn(t))
            if x is not None: tcPr.remove(x)
        tcPr.append(el("w:tcW", **{"w:w": TW(width_mm), "w:type": "dxa"}))
        b = el("w:tcBorders")
        for side, (sz, col) in borders.items():
            b.append(el("w:" + side, **{"w:val": "single" if sz else "nil", "w:sz": sz or 0, "w:space": 0, "w:color": col or "auto"}))
        tcPr.append(b)
        if fill: tcPr.append(el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
        m = el("w:tcMar")
        for side, v in (("top", vpad), ("left", lr), ("bottom", vpad), ("right", lr)):
            m.append(el("w:" + side, **{"w:w": TW(v), "w:type": "dxa"}))
        tcPr.append(m); tcPr.append(el("w:vAlign", **{"w:val": "top"}))

    def make_table(rows, widths, indent, header=True, label_col=False, vpad=3.6):
        ncol = len(widths)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        t = doc.add_table(rows=len(rows), cols=ncol); t.autofit = False
        tblPr = t._tbl.tblPr
        for x in tblPr.findall(qn("w:tblStyle")): tblPr.remove(x)
        tblPr.append(el("w:tblW", **{"w:w": TW(sum(widths)), "w:type": "dxa"}))
        tblPr.append(el("w:tblInd", **{"w:w": TW(indent), "w:type": "dxa"}))
        b = el("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"): b.append(el("w:" + side, **{"w:val": "nil"}))
        tblPr.append(b); tblPr.append(el("w:tblLayout", **{"w:type": "fixed"}))
        tblPr.append(el("w:tblLook", **{"w:val": "0000", "w:firstRow": 0, "w:lastRow": 0, "w:firstColumn": 0,
                                        "w:lastColumn": 0, "w:noHBand": 1, "w:noVBand": 1}))
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(TW(w)))
        nums = [numeric_col(rows, c, header) for c in range(ncol)]
        short_first = not label_col and widths[0] <= 20
        for ri, row in enumerate(rows):
            trPr = t.rows[ri]._tr.get_or_add_trPr()
            trPr.append(el("w:cantSplit"))                      # las filas no se parten
            is_hdr = header and ri == 0
            if is_hdr: trPr.append(el("w:tblHeader"))           # cabecera repetida en cada página
            bi = ri - (1 if header else 0)
            fill = PALE if is_hdr else (ALT if bi % 2 == 1 else None)
            for ci, txt in enumerate(row):
                c = t.cell(ri, ci)
                bd = {"top": (12, OLIVE) if ri == 0 else (0, None), "bottom": (6, SEP) if is_hdr else (4, SEP2),
                      "left": (0, None), "right": (0, None)}
                cell_set(c, widths[ci], fill=fill, vpad=3.2 if is_hdr else vpad, lr=1.6 if ci == 0 else 2.5, borders=bd)
                p = c.paragraphs[0]
                pfmt(p, line=1.0 if is_hdr else 1.12,
                     align=WD_ALIGN_PARAGRAPH.RIGHT if nums[ci] else None)
                if is_hdr: r = p.add_run(txt); set_font(r, F_MED, 9, OLIVE)
                elif label_col and ci == 0: add_runs(p, txt.replace("**", ""), 9.5, color=INK, bold_color=INK); [set_font(x, F_MED) for x in p.runs]
                elif short_first and ci == 0: add_runs(p, txt, 9.5, color=OLIVE)
                else: add_runs(p, txt, 9.5)

    def spacer():
        sp = doc.add_paragraph(); pfmt(sp, line=1.0)
        mr = el("w:rPr"); mr.append(el("w:sz", **{"w:val": 10})); mr.append(el("w:szCs", **{"w:val": 10}))
        sp._p.get_or_add_pPr().append(mr)

    if meta:   # ficha inicial: tabla de dos columnas, etiqueta compacta
        lw = min(55.0, max(30.0, max(len(k) for k, _ in meta) * 1.9 + CELL_PAD))
        make_table([[k, v] for k, v in meta], [lw, TEXT_W - lw], 0, header=False, label_col=True, vpad=2.5)

    bullets = {}
    def bullet_id():
        if bullets: return 90
        numbering = doc.part.numbering_part.element
        absn = parse_xml(
            '<w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="90">'
            '<w:multiLevelType w:val="singleLevel"/><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            '<w:lvlText w:val="•"/><w:lvlJc w:val="left"/><w:pPr><w:ind w:left="%d" w:hanging="%d"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s" w:cs="%s"/><w:color w:val="%s"/></w:rPr></w:lvl></w:abstractNum>'
            % (TW(COL + 5), TW(5), F_REG, F_REG, F_REG, OLIVE))
        first = numbering.find(qn("w:num"))
        (first.addprevious(absn) if first is not None else numbering.append(absn))
        numbering.append(parse_xml('<w:num xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                                   'w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
        bullets[1] = True; return 90

    front = bool(meta) if P["salto_primer_capitulo"] == "auto" else P["salto_primer_capitulo"].lower() in ("si", "sí", "true", "1")
    prev = None; h1_idx = 0
    for k, (kind, val) in enumerate(blocks):
        nxt = blocks[k + 1][0] if k + 1 < len(blocks) else None
        if kind == "h1":
            h1_idx += 1
            m = re.match(r"^(\d+(?:\.\d+)*\.?)\s+(.*)$", val)
            num, title = (m.group(1), m.group(2)) if m else ("", val)
            p = doc.add_paragraph(style="Heading 1")
            first_front = prev is None and front           # tras ficha o índice: capítulo en página nueva
            at_top = first_front or h1_idx in top_h1
            pfmt(p, before=0 if first_front else 30, after=16, left=COL, hanging=COL, keep_next=True)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(COL))
            if first_front: p.paragraph_format.page_break_before = True
            if not at_top: border(p, "top", SEP, sz=6, space=17)
            if num:
                r = p.add_run(num); set_font(r, F_REG, 26, OLIVE)
                r._element.get_or_add_rPr().append(el("w:spacing", **{"w:val": -20}))
                r = p.add_run(" \t" + title)
            else:
                r = p.add_run("\t" + title)
            set_font(r, F_MED, 16, OLIVE)
            r._element.get_or_add_rPr().append(el("w:position", **{"w:val": 6}))
        elif kind == "h2":
            p = doc.add_paragraph(style="Heading 2")
            pfmt(p, before=8 if prev == "h1" else 14, after=7, left=COL, right=PROSE_R, keep_next=True, line=1.1)
            r = p.add_run(val); set_font(r, F_MED, 12, INK)
        elif kind == "p":
            p = doc.add_paragraph()
            lead_only = val.startswith("**") and val.endswith("**") and val.count("**") == 2
            pfmt(p, after=5 if lead_only else 9, line=LINE, left=COL, right=PROSE_R,
                 keep_next=(nxt in ("table", "li") or lead_only))
            add_runs(p, val, 10)
        elif kind == "li":
            p = doc.add_paragraph()
            pfmt(p, after=6 if nxt == "li" else 10, line=LINE, left=COL + 5, right=PROSE_R, hanging=5)
            numPr = el("w:numPr"); numPr.append(el("w:ilvl", **{"w:val": 0})); numPr.append(el("w:numId", **{"w:val": bullet_id()}))
            p._p.get_or_add_pPr().insert(0, numPr)
            add_runs(p, val, 10)
        elif kind == "table":
            W = TEXT_W - COL
            widths = val["widths"] or auto_widths(val["rows"], W, val["header"])
            if val["widths"]: widths = [w * W / sum(widths) for w in widths]
            make_table(val["rows"], widths, COL, header=val["header"])
            print("  tabla %d col: %s mm" % (len(widths), " / ".join("%.1f" % w for w in widths)))
            spacer()
        prev = kind

    cp = doc.core_properties
    cp.title = " ".join(P["titulo"]); cp.author = "Azolla Projects"; cp.language = "es-ES"
    doc.save(out)

# ---------------- render y detección de capítulos en cabeza de página
def ensure_fonts():
    have = subprocess.run(["fc-list"], capture_output=True, text=True).stdout
    if "IBM Plex Sans" in have and "SemiBold" in have: return
    d = os.path.expanduser("~/.fonts"); os.makedirs(d, exist_ok=True)
    for f in os.listdir(os.path.join(ASSETS, "fonts")):
        shutil.copy(os.path.join(ASSETS, "fonts", f), d)
    subprocess.run(["fc-cache", "-f"], capture_output=True)

def render(docx, pdf):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, docx], capture_output=True, timeout=240)
    got = os.path.join(tmp, os.path.splitext(os.path.basename(docx))[0] + ".pdf")
    if not os.path.exists(got): return False
    shutil.move(got, pdf); return True

def h1_at_top(pdf):
    """Índices (1..n) de los capítulos cuyo encabezado es lo primero del cuerpo de su página."""
    x = subprocess.run(["pdftotext", "-bbox", pdf, "-"], capture_output=True, text=True).stdout
    tops, idx = set(), 0
    for H, body in re.findall(r'<page width="[\d.]+" height="([\d.]+)">(.*?)</page>', x, re.S):
        mm = float(H) / 297
        words = [(float(b), float(d) - float(b)) for b, d in
                 re.findall(r'<word xMin="[\d.]+" yMin="([\d.]+)" xMax="[\d.]+" yMax="([\d.]+)">', body)]
        words = [w for w in words if 32 * mm < w[0] < 268 * mm]
        if not words: continue
        first = min(w[0] for w in words)
        # encabezados de capítulo: número a 26 pt y título a 16 pt (alto de caja 19-39 pt);
        # el título del informe (30 pt, ~42) y los subapartados (12 pt, ~17) quedan fuera
        ys = sorted(w[0] for w in words if 19 < w[1] < 39)
        clusters = []
        for y in ys:
            if not clusters or y - clusters[-1][-1] > 8: clusters.append([y])
            else: clusters[-1].append(y)
        for c in clusters:
            idx += 1
            if min(c) <= first + 8: tops.add(idx)
    return tops

def doc_words(docx):
    d = Document(docx); words = []
    def ptext(p):
        out = []
        for e in p.iter():
            tag = e.tag.split("}")[1]
            if tag == "t": out.append(e.text or "")
            elif tag in ("tab", "br"): out.append(" ")
        return "".join(out)
    for blk in d.element.body:
        tag = blk.tag.split("}")[1]
        if tag == "p": words += ptext(blk).split()
        elif tag == "tbl":
            for tc in blk.iter():
                if tc.tag.endswith("}tc"):
                    words += " ".join(ptext(p) for p in tc.iter() if p.tag.endswith("}p")).split()
    return words

def verify(src, docx):
    P, _, _ = parse(src)
    expected = " ".join(P["titulo"] + [P["para"], P["fecha"]]).split() + plain_words(src)
    got = doc_words(docx)
    if got == expected: return True
    for i, (a, b) in enumerate(zip(got, expected)):
        if a != b:
            print("DIFERENCIA en palabra %d: docx=%r contenido=%r · contexto: %s" % (i, a, b, " ".join(expected[max(0, i-6):i+6])))
            return False
    print("DIFERENCIA de longitud: docx=%d palabras, contenido=%d" % (len(got), len(expected)))
    return False

def main():
    ap = argparse.ArgumentParser()
    if len(sys.argv) == 3 and sys.argv[1] == "--extraer-helecho":
        import base64
        html = open(sys.argv[2], encoding="utf-8").read()
        b = re.search(r'class="fern" src="data:image/png;base64,([A-Za-z0-9+/=]+)"', html).group(1)
        os.makedirs(ASSETS, exist_ok=True); open(os.path.join(ASSETS, "helecho.png"), "wb").write(base64.b64decode(b))
        print("Helecho extraído en", os.path.join(ASSETS, "helecho.png")); return
    ap.add_argument("contenido"); ap.add_argument("salida")
    ap.add_argument("--pdf", help="guardar también el PDF de control en esta ruta")
    ap.add_argument("--sin-render", action="store_true", help="no renderizar (sin LibreOffice)")
    a = ap.parse_args()
    ensure_assets()
    top = frozenset()
    build(a.contenido, a.salida, top)
    pdf = os.path.join(tempfile.gettempdir(), "doc_informe_render.pdf")
    rendered = False
    if not a.sin_render and shutil.which("soffice"):
        ensure_fonts()
        for it in range(4):
            if not render(a.salida, pdf): break
            rendered = True
            new = frozenset(h1_at_top(pdf))
            if new == top: break
            top = new
            build(a.contenido, a.salida, top)
        print("Capítulos en cabeza de página (sin línea superior):", sorted(top) or "ninguno")
    else:
        print("AVISO: sin render; no se ha comprobado la doble línea en cabeza de página.")
    if a.pdf and rendered: shutil.copy(pdf, a.pdf); print("PDF de control:", a.pdf)
    print("VERIFICACIÓN OK: texto idéntico al contenido" if verify(a.contenido, a.salida)
          else "VERIFICACIÓN FALLIDA: revisar antes de entregar")
    if rendered: print("Render de revisión:", pdf)
    print("OK ->", a.salida)

if __name__ == "__main__":
    main()
